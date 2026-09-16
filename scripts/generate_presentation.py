import logging
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s"
)
logger = logging.getLogger(__name__)

BASE_DIR = Path(__file__).resolve().parent.parent
REPORTS_DIR = BASE_DIR / "reports"
REPORTS_DIR.mkdir(parents=True, exist_ok=True)
CHARTS_DIR = BASE_DIR / "charts"

PPTX_PATH = REPORTS_DIR / "Bluestock_MF_Presentation.pptx"

def add_slide(prs, title_text, content_bullets, image_path=None):
    slide_layout = prs.slide_layouts[6] # Blank layout
    slide = prs.slides.add_slide(slide_layout)
    
    # Title Header Box
    tb_title = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9), Inches(0.8))
    tf_title = tb_title.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(24)
    p_title.font.bold = True
    p_title.font.color.rgb = RGBColor(30, 58, 138) # Dark Navy
    
    # Content Box
    width = Inches(5.0) if image_path and Path(image_path).exists() else Inches(8.8)
    tb_content = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), width, Inches(5.2))
    tf_content = tb_content.text_frame
    tf_content.word_wrap = True
    
    for i, bullet in enumerate(content_bullets):
        p = tf_content.add_paragraph() if i > 0 else tf_content.paragraphs[0]
        p.text = bullet
        p.font.size = Pt(15)
        p.font.color.rgb = RGBColor(31, 41, 55)
        p.space_after = Pt(10)
        
    # Image if provided
    if image_path and Path(image_path).exists():
        slide.shapes.add_picture(str(image_path), Inches(5.7), Inches(1.3), width=Inches(3.8))
        
    return slide

def generate_pptx():
    logger.info(f"Generating 12-slide PowerPoint presentation at {PPTX_PATH}...")
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625) # 16:9 widescreen
    
    # Slide 1: Title
    slide_layout = prs.slide_layouts[6]
    slide1 = prs.slides.add_slide(slide_layout)
    tb = slide1.shapes.add_textbox(Inches(1), Inches(1.5), Inches(8), Inches(2.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "BLUESTOCK FINTECH"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = RGBColor(79, 70, 229)
    
    p2 = tf.add_paragraph()
    p2.text = "Mutual Fund Analytics Platform"
    p2.font.size = Pt(32)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(30, 58, 138)
    
    p3 = tf.add_paragraph()
    p3.text = "Comprehensive 7-Day Capstone Presentation Deck"
    p3.font.size = Pt(16)
    p3.font.color.rgb = RGBColor(107, 114, 128)
    
    # Slide 2: Problem & Objective
    add_slide(prs, "1. Problem & Objectives", [
        "• Mutual fund investors face fragmented data across fund houses and benchmarks.",
        "• Difficulty analyzing risk-adjusted returns (Sharpe, Sortino, Alpha, Beta).",
        "• Objective: Build an end-to-end analytics platform from raw datasets to interactive dashboards.",
        "• Implement automated ETL, star schema database, risk analytics, and fund recommendation."
    ])
    
    # Slide 3: Data Sources
    add_slide(prs, "2. Data Sources & Authenticity", [
        "• Integrated 10 core mutual fund datasets (87,000+ total records).",
        "• Anchored to public data: Fund master, NAV history, AUM, SIP growth, portfolio holdings.",
        "• Investor transaction data synthetically generated for privacy compliance.",
        "• Data quality validation ensures 0 missing NAVs and standardized schemas."
    ])
    
    # Slide 4: Architecture
    add_slide(prs, "3. System Architecture", [
        "• Data Ingestion: Raw CSV extraction & Live NAV API fetcher (`mfapi.in`).",
        "• Data Cleaning: Standardization, duplicate removal, forward-filling NAVs.",
        "• Data Storage: SQLite Star Schema DB (`dim_fund`, `dim_date`, 8 fact tables).",
        "• Analytics & BI: Python (`Pandas`, `NumPy`, `SciPy`), Streamlit Dashboard, ReportLab PDF."
    ])
    
    # Slide 5: EDA Highlights
    add_slide(prs, "4. EDA Highlights", [
        "• Industry AUM exceeds ₹45 Lakh Crore.",
        "• Monthly SIP inflows show constant positive trajectory exceeding ₹19,000 Cr.",
        "• Folio count growth driven primarily by retail equity investments.",
        "• Top 5 fund houses dominate over 65% of total industry AUM."
    ], CHARTS_DIR / "sip_inflow_trend.png")
    
    # Slide 6: EDA Insights
    add_slide(prs, "5. Category & Sector Insights", [
        "• Equity & Small Cap funds yielded highest 3-Year CAGR returns (18% - 24%).",
        "• Banking, Financial Services, IT, and Oil & Gas represent top sector holdings.",
        "• Sector Concentration HHI indicates well-diversified portfolios across major funds."
    ], CHARTS_DIR / "sector_allocation.png")
    
    # Slide 7: Performance Metrics
    add_slide(prs, "6. Performance Analytics & Scorecard", [
        "• Calculated 1Yr, 3Yr CAGR, 5Yr CAGR, Sharpe, and Sortino ratios.",
        "• Developed Composite Fund Scorecard (0-100 score).",
        "• Weights: 30% 3Yr Return, 25% Sharpe, 20% Alpha, 15% Expense Ratio (inv), 10% Max DD (inv).",
        "• Identified top-performing funds across equity, debt, and hybrid categories."
    ], CHARTS_DIR / "top_10_funds.png")
    
    # Slide 8: Risk & Benchmark Analysis
    add_slide(prs, "7. Risk & Benchmark Analysis", [
        "• Annualized Volatility (Std Dev) ranges from 11% to 22%.",
        "• Historical 95% VaR ranges between 1.1% and 2.4% daily.",
        "• Alpha analysis shows top funds outperforming Nifty 100 benchmark by +2.4% to +6.8%.",
        "• Maximum Drawdown analysis highlights resilience during market pullbacks."
    ], CHARTS_DIR / "risk_return_relationship.png")
    
    # Slide 9: Dashboard Overview
    add_slide(prs, "8. Interactive Streamlit Dashboard", [
        "• Developed 4-page interactive Streamlit web dashboard (`dashboard/app.py`).",
        "• Page 1: Industry Overview & Macro KPIs.",
        "• Page 2: Fund Performance, Risk Scatter & Composite Scorecard.",
        "• Page 3: Investor Analytics & Geographic Splits.",
        "• Page 4: SIP & Market Trends with Benchmark Index comparisons."
    ])
    
    # Slide 10: Investor Analytics
    add_slide(prs, "9. Investor Demographics & Behaviors", [
        "• T30 cities account for 62% of investment volume; B30 locations expanding rapidly.",
        "• Investors aged 26-35 lead new monthly SIP account creation.",
        "• SIP continuity analysis flagged at-risk investors with transaction gap > 35 days.",
        "• Investor cohort retention remains strong at 78% for recent 2022-2024 cohorts."
    ], CHARTS_DIR / "geographic_transactions.png")
    
    # Slide 11: Key Findings & Recommendations
    add_slide(prs, "10. Key Findings & Strategic Recommendations", [
        "1. Promote High-Sharpe, low-cost Direct plan schemes to retail investors.",
        "2. Expand digital distribution in B30 cities to accelerate retail growth.",
        "3. Deploy automated SIP gap notifications to minimize investor drop-offs.",
        "4. Leverage transparent rule-based recommendation model for investor guidance."
    ])
    
    # Slide 12: Conclusion / Thank You
    add_slide(prs, "11. Conclusion & Deliverables", [
        "• Completed submission-ready 7-Day Mutual Fund Analytics Platform.",
        "• Integrated ETL pipelines, SQLite star schema DB, Python analytics, Streamlit dashboard.",
        "• Generated automated 24-section PDF Final Report and PowerPoint deck.",
        "• Thank you! Questions & Discussion."
    ])
    
    prs.save(str(PPTX_PATH))
    logger.info(f"PowerPoint presentation successfully saved at {PPTX_PATH}")

if __name__ == "__main__":
    generate_pptx()
