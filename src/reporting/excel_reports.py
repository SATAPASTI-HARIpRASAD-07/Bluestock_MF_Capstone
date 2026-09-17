"""
Excel Report Generation Module using openpyxl with Filters, Formatting, and Freeze Panes.
"""

import os
import sqlite3
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class ExcelReportGenerator:
    def __init__(self, db_path: str = "db/nifty100.db", output_dir: str = "outputs"):
        self.db_path = db_path
        self.output_dir = output_dir
        os.makedirs(self.output_dir, exist_ok=True)

    def generate_screener_excel(self, df: pd.DataFrame, filename: str = "screener_results.xlsx") -> str:
        """
        Export screener results to formatted Excel workbook.
        """
        file_path = os.path.join(self.output_dir, filename)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Screener Results"

        # Apply header styling
        headers = list(df.columns)
        ws.append(headers)

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")

        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Append data rows
        for _, row in df.iterrows():
            ws.append(list(row.values))

        # Formatting & Freeze Panes
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        # Auto column width
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        wb.save(file_path)
        return file_path
